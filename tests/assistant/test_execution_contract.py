"""Regression cases discovered with real English/Hinglish multi-tool runs."""

import json
from datetime import datetime, timedelta
from itertools import count
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import httpx
import pytest
from openpyxl import load_workbook

from skyview.assistant import (
    artifacts,
    engine,
    models,
    provider,
    store,
)
from skyview.assistant import (
    execution_contract as contract,
)
from skyview.assistant.tools import Executor, schemas


def test_complex_request_keeps_all_workflows_and_file_formats():
    goal = contract.infer(
        "Make a budget: 2 bags at INR 1250 each, total cost chart. Create Excel, Word and PDF. Do not save any expense."
    )
    assert set(goal["formats"]) == {"xlsx", "docx", "pdf"}
    assert "calculate_budget" in goal["reads"]
    names = {
        s["function"]["name"]
        for s in schemas(
            ["documents", "field-check", "irrigation", "market-prices"],
            pinned_skills=goal["skills"],
        )
    }
    assert {"create_file", "calculate_budget", "market_trend"} <= names
    assert (
        contract.infer("Do not create any PDF. Explain what PDF means.")["formats"]
        == []
    )
    assert contract.infer("What can you do with Excel and Word?")["formats"] == []


def test_colloquial_requests_find_work_without_tool_names():
    plain = contract.infer("Mere Nashik wale khet ka kaam aur kharcha dekh ke batao.")
    assert {"list_tasks", "list_expenses"} <= set(plain["reads"])
    goal = contract.infer(
        "Mere kaunse kaam pending hain aur ab tak kitna kharcha hua? "
        "Gehu ka mandi bhav bhi batao, aur agle winter mein kya ugau?"
    )
    assert {
        "list_tasks",
        "list_expenses",
        "market_prices",
        "seasonal_planning_context",
    } <= set(goal["reads"])
    assert {"farm-records", "market-prices", "crop-calendar"} <= set(goal["skills"])
    goal = contract.infer(
        "Seeds ke 2 bags 1250 each, compost ke 8 bags 300 each, truck 700. "
        "Hisab lagao aur put it into a sheet and printable report. Koi kharcha record mat karna."
    )
    assert set(goal["formats"]) == {"xlsx", "pdf"}
    assert "farm-budget" in goal["skills"]
    assert "save_task" not in goal["reads"]
    assert (
        "crop-scouting"
        in contract.infer(
            "My tomato leaves are turning yellow. What should I check first?"
        )["skills"]
    )


def test_receipts_ignore_old_files_skipped_calls_and_failed_artifacts():
    goal = {"reads": ["list_tasks"], "formats": ["pdf", "xlsx"], "start": 1}
    messages = [
        {
            "role": "tool",
            "content": json.dumps({"artifact": {"format": "pdf", "file_id": "old"}}),
        },
        {
            "role": "assistant",
            "tool_calls": [
                {"id": "a", "function": {"name": "list_tasks"}},
                {"id": "b", "function": {"name": "create_file"}},
            ],
        },
        {"role": "tool", "tool_call_id": "a", "content": '{"skipped":true}'},
        {"role": "tool", "tool_call_id": "b", "content": '{"error":"failed"}'},
    ]
    assert contract.remaining(goal, messages) == (["list_tasks"], ["pdf", "xlsx"])
    messages.extend(
        [
            {"role": "tool", "tool_call_id": "a", "content": '{"tasks":[]}'},
            {
                "role": "tool",
                "tool_call_id": "b",
                "content": '{"artifact":{"format":"xlsx","file_id":"new"}}',
            },
        ]
    )
    assert contract.remaining(goal, messages) == ([], ["pdf"])


def test_hinglish_persists_across_followup_but_explicit_english_wins():
    messages = [
        {"role": "user", "content": "Hinglish mein Roman letters mein bolo."},
        {"role": "user", "content": "Ab compost 8 bags kar do."},
    ]
    assert "Latin/Roman" in contract.language_instruction(messages)
    messages.append({"role": "user", "content": "Reply in English now."})
    assert "Reply in English" in contract.language_instruction(messages)


@pytest.mark.parametrize(
    ("message", "locale", "expected"),
    [
        ("मेरे खेत में गेहूँ कब बोना चाहिए?", "en", "Hindi using Devanagari"),
        ("माझ्या शेतात काय पेरू?", "mr", "Marathi using Devanagari"),
        ("আমার ধানের কী যত্ন দরকার?", "en", "Bengali"),
        ("Mere khet ke liye kya ugau?", "en", "Hinglish using Latin/Roman"),
        ("What should I grow?", "hi", "Reply in English"),
    ],
)
def test_reply_language_follows_farmers_current_words(message, locale, expected):
    current = [{"role": "user", "content": message}]
    assert expected in contract.language_instruction(current, ui_language=locale)


def test_reply_language_is_scoped_to_current_request_and_its_followup():
    messages = [
        {"role": "user", "content": "Hinglish mein batao"},
        {"role": "assistant", "content": "Theek hai"},
        {"role": "user", "content": "What should I grow?"},
    ]
    assert "Reply in English" in contract.language_instruction(messages, start=2)
    messages.append({"role": "user", "content": "March se shuru karna hai"})
    assert "Hinglish" in contract.language_instruction(messages, start=2)


def test_generated_files_never_link_to_invented_provider_urls():
    messages = [
        {
            "role": "tool",
            "content": json.dumps(
                {"artifact": {"file_id": "abc123", "filename": "Budget.xlsx"}}
            ),
        }
    ]
    text = '[Budget.xlsx](https://fake.example/abc123) and [source](https://icar.gov.in/)\n```json\n{"title":"Cost", "points":[]}\n```'
    result = contract.public_text(text, messages)
    assert "fake.example" not in result and "points" not in result
    assert "[source](https://icar.gov.in/)" in result
    assert "list_tasks" not in contract.public_text("I checked list_tasks.", [])


@pytest.mark.asyncio
async def test_prose_cannot_finish_an_unexecuted_file_request(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "ROOT", tmp_path)
    sid = store.create_session()["id"]
    calls = []

    async def lying_model(*args, **kwargs):
        calls.append(kwargs)
        yield {"type": "delta", "text": "Done! Your PDF is ready."}
        yield {
            "type": "turn",
            "message": {"role": "assistant", "content": "Done! Your PDF is ready."},
        }

    monkeypatch.setattr(provider, "stream_turn", lying_model)
    events = [
        e
        async for e in engine.run(
            sid,
            models.RunRequest(
                message="Create a PDF with a checklist: inspect leaves, check water."
            ),
            store.load(sid),
        )
    ]
    assert len(calls) == 2 and all(c["tool_choice"] == "required" for c in calls)
    assert not any(e["type"] == "delta" for e in events)
    assert events[-1]["type"] == "error"
    assert store.load(sid)["display_history"][-1]["status"] == "error"


def test_provider_schema_removes_decoding_limits_not_properties():
    original = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "minimum": {"type": "number", "minimum": 1},
            "rows": {
                "type": "array",
                "maxItems": 500,
                "items": {"type": "string", "maxLength": 4000},
            },
        },
    }
    compact = provider.serving_schema(original)
    assert compact["properties"]["minimum"] == {"type": "number"}
    assert "maxItems" not in compact["properties"]["rows"]
    assert original["properties"]["rows"]["maxItems"] == 500
    with pytest.raises(ValueError):
        models.FileArgs(
            format="xlsx", title="Too many", columns=["A"], rows=[[1]] * 501
        )


@pytest.mark.asyncio
async def test_rate_limited_credential_not_retried_each_tool_round(monkeypatch):
    monkeypatch.setattr(provider, "_cooldowns", {})
    monkeypatch.setattr(
        provider,
        "get_settings",
        lambda: SimpleNamespace(GROQ_API_KEYS=["test-only"], GEMINI_API_KEY=None),
    )
    requests = []

    def limited(request):
        requests.append(request)
        return httpx.Response(
            429, headers={"Retry-After": "30"}, json={"error": {"message": "quota"}}
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(limited)) as client:
        for _ in range(2):
            with pytest.raises(provider.ProviderUnavailable):
                _ = [
                    e async for e in provider.stream_turn(client, [], [], quota_wait=0)
                ]
    assert len(requests) == 1


@pytest.mark.asyncio
async def test_all_configured_keys_receive_balanced_first_attempts(monkeypatch):
    keys = [f"test-key-{index}" for index in range(12)]
    monkeypatch.setattr(provider, "_cooldowns", {})
    monkeypatch.setattr(provider, "_key_rotation", count())
    monkeypatch.setattr(
        provider,
        "get_settings",
        lambda: SimpleNamespace(GROQ_API_KEYS=keys, GEMINI_API_KEY=None),
    )
    first_attempts = []

    def answer(request):
        first_attempts.append(request.headers["authorization"].removeprefix("Bearer "))
        return httpx.Response(
            200,
            text='data: {"choices":[{"delta":{"content":"ok"}}]}\n\ndata: [DONE]\n\n',
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(answer)) as client:

        async def one_turn():
            return [part async for part in provider.stream_turn(client, [], [])]

        import asyncio

        results = await asyncio.gather(*(one_turn() for _ in range(24)))
    assert all(parts[-1]["type"] == "turn" for parts in results)
    assert {key: first_attempts.count(key) for key in keys} == dict.fromkeys(keys, 2)


@pytest.mark.asyncio
async def test_organization_limit_does_not_retry_every_key(monkeypatch):
    keys = [f"test-key-{index}" for index in range(12)]
    monkeypatch.setattr(provider, "_cooldowns", {})
    monkeypatch.setattr(provider, "_key_rotation", count())
    monkeypatch.setattr(
        provider,
        "get_settings",
        lambda: SimpleNamespace(GROQ_API_KEYS=keys, GEMINI_API_KEY=None),
    )
    attempts = []

    def limited(request):
        attempts.append(request.headers["authorization"])
        return httpx.Response(
            429,
            json={"error": {"message": "Rate limit reached in organization on TPM"}},
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(limited)) as client:
        with pytest.raises(provider.ProviderUnavailable):
            _ = [
                part
                async for part in provider.stream_turn(client, [], [], quota_wait=0)
            ]
    assert len(attempts) == 1
    assert provider.shared_quota_limit(httpx.Response(413))


def test_xlsx_summary_row_is_not_double_counted_or_charted(monkeypatch, tmp_path):
    monkeypatch.setattr(store, "ROOT", tmp_path)
    sid = store.create_session()["id"]
    file = artifacts.create_file(
        sid,
        models.FileArgs(
            format="xlsx",
            title="Season budget",
            columns=["Item", "Total cost"],
            rows=[
                ["Seeds", 2500],
                ["Compost", 3000],
                ["Transport", 500],
                ["Total", 6000],
            ],
        ),
    )
    path = store.ROOT / "files" / (file["file_id"] + ".xlsx")
    workbook = load_workbook(path, data_only=True)
    assert workbook["Farm records"]["B7"].value == 6000
    workbook.close()
    workbook = load_workbook(path)
    assert (
        workbook["Farm records"]._charts[0].series[0].val.numRef.f
        == "'Farm records'!$B$2:$B$4"
    )
    workbook.close()


@pytest.mark.asyncio
async def test_trend_obeys_requested_days_and_excludes_future(monkeypatch):
    from skyview.data.v3 import db

    monkeypatch.setattr(db, "enabled", lambda: False)
    today = datetime.now(ZoneInfo("Asia/Kolkata")).date()
    rows = [
        {
            "market": "One",
            "variety": "FAQ",
            "unit": "INR/quintal",
            "arrival_date": (today - timedelta(days=i)).isoformat(),
            "modal_price": 2000 + i,
        }
        for i in range(-2, 101)
    ]

    async def upstream(*args, **kwargs):
        return {"records": rows}

    async with httpx.AsyncClient() as client:
        executor = Executor("fixture", "WS01", client)
        monkeypatch.setattr(executor, "upstream", upstream)
        result = await executor.run("market_trend", {"commodity": "Wheat", "days": 90})
    assert result["total_observations"] == 90
    assert result["coverage"]["start"] == (today - timedelta(days=89)).isoformat()
    assert result["coverage"]["end"] == today.isoformat()


def test_task_creation_requires_matching_action_object():
    goal = contract.infer(
        "Create Excel, Word and PDF versions. Do not record any expense or calendar task."
    )
    assert goal["formats"] == ["xlsx", "docx", "pdf"]
    assert "save_task" not in goal["reads"]
    goal = contract.infer(
        "Create one calendar task named Inspect pipes. Do not create any other task."
    )
    assert goal["reads"] == ["save_task"]
    assert (
        "save_task"
        not in contract.infer("Do not save any task. Explain the calendar.")["reads"]
    )
    assert contract.infer("Move only Inspect pipes to 2026-11-07.")["reads"] == [
        "list_tasks",
        "update_task",
    ]


def test_provider_daily_quota_reset_is_not_shortened_to_minute():
    response = httpx.Response(
        429, json=[{"error": {"details": [{"retryDelay": "23400.5s"}]}}]
    )
    assert provider.quota_delay(response) == 23400.5


def test_input_file_mentions_are_not_mistaken_for_created_deliverables():
    goal = contract.infer("Read the attached CSV and create a PDF summary.")
    assert goal["formats"] == ["pdf"]
    assert contract.infer(
        "Draft a negotiation note and a Word checklist for handover."
    )["formats"] == ["docx"]
    assert not contract.unverified_file_claims(
        "I read **budget.csv** and made **Budget.pdf**.",
        [{"filename": "Budget.pdf"}],
        known_names=["budget.csv"],
    )
    assert contract.unverified_file_claims(
        "Created **Another.pdf**.", [{"filename": "Budget.pdf"}]
    ) == ["Another.pdf"]


def test_model_chart_excerpt_does_not_change_public_tool_result():
    source = {
        "records": [{"value": n} for n in range(100)],
        "chart": {"points": [{"date": str(n), "value": n} for n in range(365)]},
    }
    result = json.loads(engine.bounded_result(source))
    assert result["chart"]["full_point_count"] == 365
    assert len(result["chart"]["points"]) == 12
    assert result["records_total"] == 100
    assert len(source["chart"]["points"]) == 365 and len(source["records"]) == 100


def test_pdf_rupee_sign_does_not_fail_with_older_unicode_font(monkeypatch, tmp_path):
    from pypdf import PdfReader
    from reportlab.pdfbase.ttfonts import TTFont

    monkeypatch.setattr(store, "ROOT", tmp_path)

    class OldFont(TTFont):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.face.charToGlyph.pop(ord("₹"), None)

    monkeypatch.setattr("reportlab.pdfbase.ttfonts.TTFont", OldFont)
    sid = store.create_session()["id"]
    result = artifacts.create_file(
        sid,
        models.FileArgs(
            format="pdf",
            title="Budget",
            sections=[models.Section(heading="Costs", text="Seeds ₹6000")],
        ),
    )
    text = "".join(
        page.extract_text()
        for page in PdfReader(store.ROOT / "files" / (result["file_id"] + ".pdf")).pages
    )
    assert "INR" in text and "6000" in text


def test_conditional_plan_rejects_invented_chemical_and_water_doses():
    plan = {"skills": ["crop-calendar"]}
    for text in ["Apply 20 kg/acre urea.", "Drip can supply 25–30 mm/week."]:
        with pytest.raises(ValueError, match="unverified"):
            contract.validate_seasonal_draft({"sections": [{"text": text}]}, plan)
    contract.validate_seasonal_draft(
        {
            "sections": [
                {
                    "text": "Compare relative water needs; get a soil test before fertilizer planning."
                }
            ]
        },
        plan,
    )
    contract.validate_seasonal_draft(
        {"sections": [{"text": "Reported fertilizer application: 20 kg/acre urea."}]},
        {"skills": ["documents"]},
    )


def test_unsourced_seasonal_answer_rejects_precise_costs_and_water():
    goal = {"skills": ["crop-calendar"], "start": 0}
    assert contract.unsupported_seasonal_claims(
        "Chana needs 350-400 mm water and seed costs ₹30-35 per kg.", goal, []
    )
    assert contract.unsupported_seasonal_claims("The crop takes 90-100 days.", goal, [])
    assert (
        contract.unsupported_seasonal_claims(
            "With limited water, chana may be the lower-risk option. Check local seed rates.",
            goal,
            [],
        )
        == []
    )


def test_record_summary_checks_complete_ledger_and_unknown_soil():
    goal = {
        "skills": ["crop-calendar"],
        "reads": ["seasonal_planning_context"],
        "formats": [],
        "start": 0,
    }
    messages = [
        {"role": "user", "content": "Check my farm work and spending in Nashik."},
        {
            "role": "assistant",
            "tool_calls": [
                {"id": "a", "function": {"name": "list_expenses"}},
                {"id": "b", "function": {"name": "seasonal_planning_context"}},
            ],
        },
        {
            "role": "tool",
            "tool_call_id": "a",
            "content": json.dumps({"expenses": [{"amount": 100}] * 5}),
        },
        {
            "role": "tool",
            "tool_call_id": "b",
            "content": json.dumps({"farm_context": {"soil_texture": "unknown"}}),
        },
    ]
    issues = contract.unsupported_record_claims(
        "Total recorded spend: ₹300. These are the only expense rows. My loam soil is suitable.",
        goal,
        messages,
    )
    assert len(issues) == 3
    assert contract.unsupported_record_claims(
        "Ye sirf teen hain, koi aur expense abhi record nahi hai. Chana ka price wheat se kam hai.",
        goal,
        messages,
    )
    assert contract.unsupported_record_claims(
        "| Task ID | Notes |\n| fc38eb6a969b52f49d8f9f420dbd5c73 | Demo plan |",
        goal,
        messages,
    )
    assert contract.unsupported_record_claims(
        "Chana looks safer with little water. Costs may vary.", goal, messages
    )
    assert not contract.unsupported_record_claims(
        "Five expense entries total ₹500. Soil texture is unknown.", goal, messages
    )
    safe = contract.verified_summary(goal, messages, hinglish=True)
    assert "5 entries" in safe and "500.00" in safe
    assert "loam" not in safe
    unfinished = contract.verified_summary(
        {**goal, "reads": ["seasonal_planning_context", "market_prices"]},
        messages,
    )
    assert "Still unfinished: market prices" in unfinished


def test_crop_comparison_does_not_add_an_unrequested_market_lookup():
    goal = contract.infer(
        "Show a 30-day Wheat market trend. Compare chana versus wheat for the next season. Create PDF and Excel."
    )
    assert "market_trend" in goal["reads"] and "market_prices" not in goal["reads"]
    goal = contract.infer(
        "Compare wheat markets in Punjab and Haryana and show a 90-day trend."
    )
    assert {"market_trend", "market_prices"} <= set(goal["reads"])


def test_returned_market_records_block_blanket_no_data_claim():
    goal = {"reads": ["market_prices"], "formats": [], "start": 0}
    messages = [
        {
            "role": "assistant",
            "tool_calls": [{"id": "prices", "function": {"name": "market_prices"}}],
        },
        {
            "role": "tool",
            "tool_call_id": "prices",
            "content": json.dumps(
                {"records": [{"commodity": "Rice", "modal_price": 2120}]}
            ),
        },
    ]
    assert contract.unsupported_record_claims(
        "Because no dated market data is available, I cannot help.", goal, messages
    )
    assert not contract.unsupported_record_claims(
        "A Rice planning record is available; it is not a reported quote.",
        goal,
        messages,
    )
