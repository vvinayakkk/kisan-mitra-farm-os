"""Regression coverage for reviewed workspace data and file boundaries."""

import json
from datetime import UTC, datetime, timedelta

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from skyview.assistant import artifacts, models, store
from skyview.assistant.data_sources import station_snapshot
from skyview.assistant.tools import Executor


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "ROOT", tmp_path)


@pytest.mark.parametrize("format", ["docx", "pdf", "xlsx", "csv", "pptx"])
def test_preview_matches_generated_content_and_is_bounded(format):
    sid = store.create_session()["id"]
    args = models.FileArgs(
        format=format,
        title="Public sample records",
        sections=[{"heading": "Notes", "text": "Sample details " * 100}],
        columns=["Item", "Cost"],
        rows=[["Sample", n] for n in range(10)],
    )
    artifact = artifacts.create_file(sid, args)
    preview = artifact["preview"]
    assert len(preview["rows"]) == 6 and preview["row_count"] == 10
    assert preview["partial"]
    assert len(json.dumps(preview)) < 3000
    path = store.ROOT / "files" / f"{artifact['file_id']}.{format}"
    if format == "csv":
        assert preview["sections"] == []
        assert "Sample,0" in path.read_text(encoding="utf-8-sig")
    if format == "pptx":
        tables = [
            shape.table
            for slide in Presentation(path).slides
            for shape in slide.shapes
            if shape.has_table
        ]
        assert sum(len(table.rows) - 1 for table in tables) == 10
        assert tables[0].cell(1, 0).text == preview["rows"][0][0]


def test_old_files_remain_owned_and_downloadable_after_100_newer_records():
    from skyview.assistant.dev import app

    owner, other = store.create_session(), store.create_session()
    artifact = artifacts.create_file(
        owner["id"],
        models.FileArgs(
            format="csv", title="Old sample", columns=["Item"], rows=[["Sample"]]
        ),
    )
    for _ in range(105):
        store.add_record(owner["id"], "artifact", {"file_id": "placeholder"})
    client = TestClient(app)
    path = f"/api/assistant/sessions/{owner['id']}/files/{artifact['file_id']}"
    assert (
        client.get(
            path, headers={"Authorization": "Bearer " + owner["token"]}
        ).status_code
        == 200
    )
    path = path.replace(owner["id"], other["id"])
    assert (
        client.get(
            path, headers={"Authorization": "Bearer " + other["token"]}
        ).status_code
        == 404
    )


def test_deleted_records_do_not_hide_active_tasks():
    sid = store.create_session()["id"]
    active = store.add_record(sid, "task", {"title": "Active sample"})
    for _ in range(105):
        store.add_record(sid, "task", {"title": "Deleted sample", "deleted": True})
    assert store.records(sid, "task")[0]["id"] == active["id"]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "stamp", ["bad", (datetime.now(UTC) + timedelta(days=1)).isoformat()]
)
async def test_invalid_and_future_station_times_never_become_live(stamp):
    async def reader(*args):
        return {"timestamp": stamp, "temperature": 25}

    async with httpx.AsyncClient() as client:
        assert (await station_snapshot(client, "WS01", reader))["available"] is False


@pytest.mark.asyncio
async def test_station_adapter_preserves_missing_readings_and_reports_staleness():
    async def reader(*args):
        return {
            "timestamp": (datetime.now(UTC) - timedelta(minutes=2)).isoformat(),
            "temperature": 0,
            "rainfall": None,
        }

    async with httpx.AsyncClient() as client:
        result = await station_snapshot(client, "WS01", reader)
        assert result["available"] and result["stale"]
        assert result["data"]["temperature"] == 0 and result["data"]["rainfall"] is None


@pytest.mark.asyncio
async def test_photos_are_not_read_as_text():
    sid = store.create_session()["id"]
    photo = store.add_record(sid, "attachment", {"format": "png"})
    async with httpx.AsyncClient() as client:
        with pytest.raises(ValueError, match="Crop Doctor"):
            await Executor(sid, "WS01", client).run(
                "read_file", {"file_id": photo["id"]}
            )


def test_sensor_ingestion_validation_missing_values_and_db_failure(monkeypatch):
    from skyview.api import sensor_routes

    app = FastAPI()
    app.include_router(sensor_routes.router)
    client = TestClient(app)
    saved = []
    monkeypatch.setattr(
        sensor_routes, "_insert_weather", lambda params: saved.append(params) or True
    )
    response = client.post(
        "/api/sensors/ingest", json={"station_id": "TEST", "temperature": 0}
    )
    assert response.status_code == 200
    assert saved[0]["rain"] is None and saved[0]["ts"].tzinfo is UTC
    assert client.post("/api/sensors/data", json=[]).status_code == 422
    assert (
        client.post("/api/sensors/ingest", json={"station_id": "../bad"}).status_code
        == 422
    )
    assert (
        client.post("/api/sensors/ingest", json={"temperature": "NaN"}).status_code
        == 422
    )
    monkeypatch.setattr(sensor_routes, "_insert_weather", lambda params: False)
    assert (
        client.post("/api/sensors/ingest", json={"station_id": "TEST"}).status_code
        == 503
    )
    monkeypatch.setattr(sensor_routes, "get_latest_weather", lambda sid: None)
    result = client.get("/api/sensors/latest/TEST").json()
    assert result["available"] is False and "temperature" not in result


def test_json_logging_format_is_valid():
    import logging

    from skyview.utils.logger import JsonFormatter

    record = logging.LogRecord(
        "test", logging.INFO, "", 0, "Sample %s", ("message",), None
    )
    assert json.loads(JsonFormatter().format(record))["message"] == "Sample message"


@pytest.mark.asyncio
async def test_legacy_advisor_does_not_invent_missing_sensor_data(monkeypatch):
    from skyview.api import advisor_routes

    monkeypatch.setattr(advisor_routes, "get_latest_weather", lambda sid: None)
    result = await advisor_routes.advisor_insights(advisor_routes.AdvisorReq())
    assert result["status"] == "unavailable"
    assert result["sensor_snapshot"] == {} and result["ai_insights"] is None


@pytest.mark.asyncio
async def test_legacy_advisor_provider_failure_does_not_invent_suitability(monkeypatch):
    from skyview.api import advisor_routes

    monkeypatch.setattr(
        advisor_routes, "get_latest_weather", lambda sid: {"temperature": 20}
    )

    async def unavailable(*args, **kwargs):
        return None

    monkeypatch.setattr(advisor_routes, "invoke_llm", unavailable)
    result = await advisor_routes.advisor_insights(
        advisor_routes.AdvisorReq(category="crops")
    )
    assert result["status"] == "unavailable" and result["ai_insights"] is None


@pytest.mark.asyncio
async def test_production_compression_does_not_buffer_assistant_stream():
    from skyview.main import SelectiveGZipMiddleware

    sent = []

    async def app(scope, receive, send):
        await send(
            {
                "type": "http.response.start",
                "status": 200,
                "headers": [(b"content-type", b"application/x-ndjson")],
            }
        )
        await send(
            {
                "type": "http.response.body",
                "body": b'{"type":"status"}\n',
                "more_body": True,
            }
        )
        assert any(event.get("body") == b'{"type":"status"}\n' for event in sent)
        await send(
            {
                "type": "http.response.body",
                "body": b'{"type":"done"}\n',
                "more_body": False,
            }
        )

    async def receive():
        return {"type": "http.request", "body": b""}

    async def send(event):
        sent.append(event)

    middleware = SelectiveGZipMiddleware(app)
    await middleware(
        {
            "type": "http",
            "path": "/api/assistant/sessions/sample/run",
            "headers": [(b"accept-encoding", b"gzip")],
        },
        receive,
        send,
    )
    assert not any(key == b"content-encoding" for key, _ in sent[0]["headers"])


def test_workspace_state_and_station_responses_are_not_cached(monkeypatch):
    from skyview.assistant.dev import app

    client = TestClient(app)
    session = client.post("/api/assistant/sessions")
    assert session.headers["cache-control"] == "no-store"
    sid, token = session.json()["id"], session.json()["token"]
    result = client.get(
        f"/api/assistant/sessions/{sid}", headers={"Authorization": "Bearer " + token}
    )
    assert result.headers["cache-control"] == "no-store"


def test_workspace_panels_share_scoped_validated_tools(monkeypatch):
    from skyview.assistant import engine
    from skyview.assistant.dev import app

    owner, other = store.create_session(), store.create_session()
    client = TestClient(app)
    prefix = f"/api/assistant/sessions/{owner['id']}/inspect/"
    headers = {"Authorization": "Bearer " + owner["token"]}

    async def read(self, name, args, execution_id=None):
        return {
            "records": [
                {
                    "market": "Public sample",
                    "arrival_date": "2026-10-14",
                    "modal_price": 10,
                }
            ]
        }

    monkeypatch.setattr(Executor, "_execute", read)
    response = client.post(
        prefix + "market_prices", headers=headers, json={"commodity": "Wheat"}
    )
    assert (
        response.status_code == 200 and response.headers["cache-control"] == "no-store"
    )
    assert store.records(owner["id"], "observation")[0]["tool"] == "market_prices"
    assert "Public sample" in engine.system_prompt(
        models.RunRequest(workspace_view="markets"), owner["id"]
    )
    assert "Public sample" not in engine.system_prompt(models.RunRequest(), other["id"])
    assert (
        client.post(prefix + "save_task", headers=headers, json={}).status_code == 404
    )
    assert (
        client.post(
            prefix + "market_prices", headers=headers, json={"commodity": ""}
        ).status_code
        == 422
    )
    assert (
        client.post(
            prefix + "market_prices",
            headers={"Authorization": "Bearer " + other["token"]},
            json={"commodity": "Wheat"},
        ).status_code
        == 401
    )
    assert store.records(owner["id"], "task") == []


def test_workspace_panel_failure_never_creates_an_observation(monkeypatch):
    from skyview.assistant.dev import app

    owner = store.create_session()

    async def unavailable(*args, **kwargs):
        raise httpx.ConnectError("Public sample failure")

    monkeypatch.setattr(Executor, "run", unavailable)
    response = TestClient(app).post(
        f"/api/assistant/sessions/{owner['id']}/inspect/scheme_search",
        headers={"Authorization": "Bearer " + owner["token"]},
        json={},
    )
    assert response.status_code == 503
    assert store.records(owner["id"], "observation") == []


def test_encoded_upload_preserves_file_validation_and_owner_scope():
    import base64

    from skyview.assistant.dev import app

    owner, other = store.create_session(), store.create_session()
    client = TestClient(app)
    path = f"/api/assistant/sessions/{owner['id']}/attachments-json"
    headers = {"Authorization": "Bearer " + owner["token"]}
    response = client.post(
        path,
        headers=headers,
        json={
            "filename": "public.txt",
            "content": base64.b64encode(b"Public test fixture").decode(),
        },
    )
    assert response.status_code == 200
    assert "text" not in response.json()
    assert (
        store.find_record(owner["id"], "attachment", response.json()["id"])["text"]
        == "Public test fixture"
    )
    assert store.records(other["id"], "attachment") == []
    assert (
        client.post(
            path, headers=headers, json={"filename": "public.txt", "content": "!!!!"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            path,
            headers={"Authorization": "Bearer " + other["token"]},
            json={"filename": "public.txt", "content": "AAAA"},
        ).status_code
        == 401
    )
