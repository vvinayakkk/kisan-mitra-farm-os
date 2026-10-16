"""Farmer workflow tools: arithmetic, document lineage and ownership."""

import httpx
import pytest
from docx import Document
from pydantic import ValidationError

from skyview.assistant import artifacts, models, store
from skyview.assistant.tools import TOOLS, Executor


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "ROOT", tmp_path)
    monkeypatch.setenv("SKYVIEW_DATA_MODE", "legacy")


@pytest.mark.asyncio
async def test_net_realization_preserves_units_and_costs():
    async with httpx.AsyncClient() as client:
        result = await Executor(store.create_session()["id"], "WS01", client).run(
            "net_sale",
            {
                "quantity_quintals": 20,
                "price_per_quintal": 2500,
                "transport_cost": 2000,
                "handling_cost": 500,
                "commission_percent": 2,
                "other_cost": 100,
            },
        )
    assert result["gross_inr"] == 50000
    assert result["costs_inr"] == 3600
    assert result["net_inr"] == 46400
    assert result["net_per_quintal_inr"] == 2320
    assert len(result["chart"]["points"]) == 3
    with pytest.raises(ValidationError):
        models.NetSaleArgs(quantity_quintals=0, price_per_quintal=2500)
    with pytest.raises(ValidationError):
        models.NetSaleArgs(quantity_quintals=20, price_per_quintal=float("nan"))


@pytest.mark.asyncio
@pytest.mark.parametrize("format", ["docx", "pdf", "xlsx", "csv", "pptx"])
async def test_revision_preserves_original_bytes_and_owned_lineage(format):
    sid = store.create_session()["id"]
    other = store.create_session()["id"]
    original = artifacts.create_file(
        sid,
        models.FileArgs(
            format=format,
            title="Season",
            sections=[{"heading": "Plan", "text": "Original dated plan"}],
            columns=["Item", "Amount"],
            rows=[["Seeds", 100]],
        ),
    )
    original_path = store.ROOT / "files" / f"{original['file_id']}.{format}"
    before = original_path.read_bytes()
    args = {
        "file_id": original["file_id"],
        "format": format,
        "title": "Season revised",
        "sections": [{"heading": "Plan", "text": "Updated plan"}],
        "columns": ["Item", "Amount"],
        "rows": [["Seeds", 200]],
    }
    async with httpx.AsyncClient() as client:
        revision = (await Executor(sid, "WS01", client).run("revise_file", args))[
            "artifact"
        ]
        with pytest.raises(ValueError, match="Original file"):
            await Executor(other, "WS01", client).run("revise_file", args)
    assert revision["revision"] == 2 and revision["revision_of"] == original["file_id"]
    assert original_path.read_bytes() == before
    path = store.ROOT / "files" / f"{revision['file_id']}.{format}"
    if format == "docx":
        revised = Document(path)
        assert any(p.text == "Updated plan" for p in revised.paragraphs)
    elif format == "pdf":
        from pypdf import PdfReader

        assert "Updated plan" in PdfReader(path).pages[0].extract_text()
    elif format == "xlsx":
        from openpyxl import load_workbook

        assert load_workbook(path, data_only=True).worksheets[0]["B2"].value == 200
    elif format == "csv":
        assert "Seeds,200" in path.read_text()
    else:
        from pptx import Presentation

        assert any(
            "Updated plan" in shape.text
            for slide in Presentation(path).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    assert len(store.records(sid, "artifact")) == 2


def test_all_workflow_mutations_need_review_and_mission_starts_incomplete():
    for tool in (
        "record_ledger",
        "create_crop_case",
        "record_crop_note",
        "create_mission",
    ):
        assert TOOLS[tool].approval
    with pytest.raises(ValidationError):
        models.MissionArgs(
            title="Irrigation", steps=[{"title": "Inspect field", "done": True}]
        )


@pytest.mark.asyncio
async def test_document_photo_reading_is_owned_reviewed_and_does_not_save_money(
    monkeypatch,
):
    import io

    from PIL import Image

    from skyview.assistant import attachments, document_vision

    sid, other = store.create_session()["id"], store.create_session()["id"]
    output = io.BytesIO()
    Image.new("RGB", (200, 200), "white").save(output, format="PNG")
    attached = attachments.extract(sid, "receipt.png", output.getvalue())
    called = []

    async def extract(image):
        called.append(image)
        return {
            "recognized_text": "Seeds: INR 1250.00",
            "unclear_parts": ["Date unreadable"],
            "document_kind": "receipt",
            "source": "test fixture",
            "model": "fixture",
        }

    monkeypatch.setattr(document_vision, "provider_extract", extract)
    with pytest.raises(ValueError, match="owned"):
        await document_vision.read_image(other, attached["id"])
    assert not called
    result = await document_vision.read_image(sid, attached["id"])
    assert result["requires_confirmation"] and result["saved"] is False
    assert result["untrusted_content"] == "Seeds: INR 1250.00"
    assert not store.records(sid, "expense")
    with Image.open(io.BytesIO(called[0])) as img:
        assert img.format == "JPEG" and not img.getexif()


@pytest.mark.asyncio
async def test_unreadable_document_photo_never_invents_text(monkeypatch):
    import io

    from PIL import Image

    from skyview.assistant import attachments, document_vision

    sid = store.create_session()["id"]
    output = io.BytesIO()
    Image.new("RGB", (200, 200), "white").save(output, format="PNG")
    attached = attachments.extract(sid, "blank.png", output.getvalue())

    async def extract(image):
        return {
            "recognized_text": "",
            "unclear_parts": ["Blank image"],
            "document_kind": "unknown",
            "source": "test fixture",
            "model": "fixture",
        }

    monkeypatch.setattr(document_vision, "provider_extract", extract)
    result = await document_vision.read_image(sid, attached["id"])
    assert not result["available"] and result["saved"] is False
