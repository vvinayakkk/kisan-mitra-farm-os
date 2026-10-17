"""Bounded text extraction. Uploads never execute formulas, macros or code."""

import csv
import io
import json
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from . import store
from skyview.data import object_storage

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 20000
FORMATS = {".txt", ".csv", ".xlsx", ".docx", ".pdf", ".pptx"}


def extract(sid, filename, content, *, persist=True):
    name = PurePosixPath(filename.replace("\\", "/")).name[:120]
    extension = PurePosixPath(name).suffix.lower()
    if extension in (".jpg", ".jpeg", ".png", ".webp"):
        import uuid

        from PIL import Image

        if not content or len(content) > MAX_BYTES:
            raise ValueError("Crop photos must be smaller than 5 MB.")
        with Image.open(io.BytesIO(content)) as image:
            if image.width * image.height > 20_000_000:
                raise ValueError("Choose a crop photo below 20 megapixels.")
            image.verify()
            fmt = image.format
        if fmt not in ("JPEG", "PNG", "WEBP"):
            raise ValueError("Choose a JPEG, PNG or WebP crop photo.")
        suffix, mime = {
            "JPEG": ("jpg", "image/jpeg"),
            "PNG": ("png", "image/png"),
            "WEBP": ("webp", "image/webp"),
        }[fmt]
        filename = uuid.uuid4().hex + "." + suffix
        folder = store.ROOT / "photos"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / filename).write_bytes(content)
        try:
            object_storage.put(folder / filename, store.ROOT, "assistant")
            return store.add_record(
                sid,
                "attachment",
                {
                    "name": name,
                    "format": suffix,
                    "mime": mime,
                    "filename": filename,
                    "characters": 0,
                    "truncated": False,
                },
            )
        except BaseException:
            (folder / filename).unlink(missing_ok=True)
            raise
    if extension not in FORMATS:
        raise ValueError("Choose a PDF, Word, Excel, CSV or text file.")
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Files must be between 1 byte and 5 MB.")
    stream = io.BytesIO(content)
    truncated = False
    if extension in (".xlsx", ".docx", ".pptx"):
        try:
            with ZipFile(stream) as archive:
                if (
                    sum(f.file_size for f in archive.infolist()) > 25 * 1024 * 1024
                    or len(archive.infolist()) > 2000
                ):
                    raise ValueError(
                        "This document expands beyond the safe reading limit."
                    )
        except BadZipFile as exc:
            raise ValueError("This file is not a readable Office document.") from exc
        stream.seek(0)
    if extension == ".txt":
        text = content.decode("utf-8-sig")
    elif extension == ".csv":
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        rows = []
        for index, row in enumerate(reader):
            if index >= 300:
                truncated = True
                break
            truncated = truncated or len(row) > 20
            rows.append(row[:20])
        text = json.dumps(rows, ensure_ascii=False)
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(
            stream, read_only=True, data_only=True, keep_links=False
        )
        sections = []
        truncated = len(workbook.worksheets) > 5
        for sheet in workbook.worksheets[:5]:
            truncated = (
                truncated or (sheet.max_row or 0) > 200 or (sheet.max_column or 0) > 20
            )
            rows = [
                row
                for row in sheet.iter_rows(
                    max_row=min(sheet.max_row or 200, 200),
                    max_col=min(sheet.max_column or 20, 20),
                    values_only=True,
                )
                if any(value is not None for value in row)
            ]
            sections.append({"sheet": sheet.title, "rows": rows})
        workbook.close()
        text = json.dumps(sections, ensure_ascii=False, default=str)
    elif extension == ".docx":
        from docx import Document

        document = Document(stream)
        text = "\n".join(p.text for p in document.paragraphs)
        truncated = len(document.tables) > 10
        for table in document.tables[:10]:
            truncated = truncated or len(table.rows) > 100
            text += "\n" + json.dumps(
                [[cell.text for cell in row.cells] for row in table.rows[:100]],
                ensure_ascii=False,
            )
    elif extension == ".pptx":
        from pptx import Presentation

        deck = Presentation(stream)
        truncated = len(deck.slides) > 50
        text = "\n".join(
            shape.text
            for slide in list(deck.slides)[:50]
            for shape in slide.shapes
            if shape.has_text_frame
        )
    else:
        from pypdf import PdfReader

        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError("Upload an unlocked PDF.")
        if len(reader.pages) > 20:
            raise ValueError("Choose a PDF with at most 20 pages.")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        raise ValueError(
            "No readable text was found. Scanned images need OCR before upload."
        )
    data = {
        "name": name,
        "format": extension[1:],
        "text": text[:MAX_TEXT],
        "characters": min(len(text), MAX_TEXT),
        "truncated": truncated or len(text) > MAX_TEXT,
    }
    if not persist:
        return data
    from skyview.data.v3.db import enabled

    if enabled():
        import uuid

        filename = uuid.uuid4().hex + extension
        folder = store.ROOT / "uploads"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / filename
        path.write_bytes(content)
        data.update(filename=filename, size=len(content))
        try:
            object_storage.put(path, store.ROOT, "assistant")
            return store.add_record(sid, "attachment", data)
        except BaseException:
            path.unlink(missing_ok=True)
            raise
    return store.add_record(sid, "attachment", data)


def manifest(sid):
    return [
        {k: v for k, v in item.items() if k != "text"}
        for item in store.records(sid, "attachment")
    ]


def file_manifest(sid):
    return [
        {
            "id": item["id"],
            "name": item["name"],
            "format": item["format"],
            "source": "uploaded",
        }
        for item in manifest(sid)
    ] + [
        {
            "id": item["file_id"],
            "name": item["filename"],
            "format": item["format"],
            "source": "generated",
        }
        for item in store.records(sid, "artifact")
    ]
